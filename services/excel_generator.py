from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


def generate_clients_debits_excel(clients_data):
    buffer = BytesIO()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clientes e Débitos"

    sheet["A1"] = "Relatório de Clientes e Débitos"
    sheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    headers = [
        "Cliente",
        "CPF/CNPJ",
        "Descrição",
        "Valor",
        "Vencimento",
        "Situação",
    ]

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(
            row=3,
            column=column,
            value=header,
        )

        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
        )

    row = 4

    total = 0
    total_pago = 0
    total_pendente = 0

    for client, debits in clients_data:
        if not debits:
            sheet.cell(
                row=row,
                column=1,
                value=client.name,
            )

            sheet.cell(
                row=row,
                column=2,
                value=client.document,
            )

            sheet.cell(
                row=row,
                column=3,
                value="Nenhum débito",
            )

            row += 1
            continue

        for debit in debits:
            valor = float(debit.amount)

            total += valor

            if debit.paid:
                situacao = "Pago"
                total_pago += valor
            else:
                situacao = "Pendente"
                total_pendente += valor

            sheet.cell(
                row=row,
                column=1,
                value=client.name,
            )

            sheet.cell(
                row=row,
                column=2,
                value=client.document,
            )

            sheet.cell(
                row=row,
                column=3,
                value=debit.description,
            )

            sheet.cell(
                row=row,
                column=4,
                value=valor,
            )

            sheet.cell(
                row=row,
                column=5,
                value=debit.due_date,
            )

            sheet.cell(
                row=row,
                column=6,
                value=situacao,
            )

            row += 1

    for current_row in range(4, row):
        sheet.cell(
            row=current_row,
            column=4,
        ).number_format = "R$ #,##0.00"

    row += 1

    sheet.cell(
        row=row,
        column=1,
        value="Resumo",
    ).font = Font(bold=True)

    row += 1

    sheet.cell(
        row=row,
        column=1,
        value="Total",
    ).font = Font(bold=True)

    sheet.cell(
        row=row,
        column=2,
        value=total,
    )

    row += 1

    sheet.cell(
        row=row,
        column=1,
        value="Total pago",
    )

    sheet.cell(
        row=row,
        column=2,
        value=total_pago,
    )

    row += 1

    sheet.cell(
        row=row,
        column=1,
        value="Total pendente",
    )

    sheet.cell(
        row=row,
        column=2,
        value=total_pendente,
    )

    for current_row in range(row - 2, row + 1):
        sheet.cell(
            row=current_row,
            column=2,
        ).number_format = "R$ #,##0.00"

    widths = {
        "A": 30,
        "B": 22,
        "C": 35,
        "D": 18,
        "E": 18,
        "F": 18,
    }

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A4"

    sheet.auto_filter.ref = f"A3:F{row - 5}"

    workbook.save(buffer)

    buffer.seek(0)

    return buffer
